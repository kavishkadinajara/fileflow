"""
ATS Resume Optimizer — deterministic resume ↔ job-description match analysis.

Applicant Tracking Systems rank a CV by how well it matches a job description, and
people pay subscription tools (Jobscan etc.) to preview that. This does the same with
a transparent, explainable algorithm — no AI, no per-call cost — so every number can
be justified to the user (and an examiner).

Pipeline
--------
1. Extract     — CV text + layout (PDF via PyMuPDF / DOCX via python-docx); JD text.
2. Keyphrases  — RAKE (degree/frequency word scoring) pulls the JD's salient phrases.
3. Skills      — match a curated skills taxonomy (aliases normalised) in JD and CV.
4. Similarity  — TF-IDF cosine between the whole JD and CV (overall topical fit).
5. Parse-ability — structural checks an ATS parser actually trips on: multi-column
                   layout, tables, image-only pages, contact info in headers, missing
                   sections, etc. (reuses the deterministic structure engine).
6. Score       — explainable weighted blend → overall 0-100 + sub-scores + a gap
                 report (missing keywords/skills, format warnings, fixes).

The optional AI step (elsewhere) only rewrites EXISTING bullet points to weave in
missing keywords — it never invents experience, because the scoring already told it
exactly which real gaps to address.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

# ──────────────────────────────────────────────────────────────────────────────
# Skills taxonomy + alias normalisation
# ──────────────────────────────────────────────────────────────────────────────
# A curated (not exhaustive) skills list. Matching against a known vocabulary is far
# more precise than free keyphrase matching for the part that matters most — the hard
# and soft skills an ATS keys on. Grouped only for readability; matching is flat.

_SKILLS: dict[str, list[str]] = {
    "languages": [
        "python", "java", "javascript", "typescript", "c", "c++", "c#", "go", "rust",
        "kotlin", "swift", "php", "ruby", "scala", "r", "matlab", "dart", "perl",
        "objective-c", "bash", "shell", "powershell", "sql", "html", "css", "sass",
    ],
    "frameworks": [
        "react", "angular", "vue", "svelte", "next.js", "nuxt", "node", "express",
        "django", "flask", "fastapi", "spring", "spring boot", ".net", "asp.net",
        "laravel", "rails", "flutter", "react native", "tensorflow", "pytorch",
        "keras", "scikit-learn", "pandas", "numpy", "tailwind", "bootstrap", "jquery",
    ],
    "databases": [
        "mysql", "postgresql", "mongodb", "redis", "sqlite", "oracle", "sql server",
        "mariadb", "cassandra", "dynamodb", "elasticsearch", "firebase", "supabase",
        "neo4j", "snowflake",
    ],
    "cloud_devops": [
        "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "terraform",
        "ansible", "jenkins", "github actions", "gitlab ci", "ci/cd", "linux", "nginx",
        "serverless", "lambda", "ec2", "s3", "cloudformation", "prometheus", "grafana",
    ],
    "tools": [
        "git", "github", "gitlab", "bitbucket", "jira", "confluence", "figma", "postman",
        "swagger", "vs code", "intellij", "selenium", "cypress", "jest", "junit",
        "webpack", "vite", "npm", "maven", "gradle",
    ],
    "data_ml": [
        "machine learning", "deep learning", "data analysis", "data science", "nlp",
        "computer vision", "data visualization", "power bi", "tableau", "excel",
        "statistics", "etl", "spark", "hadoop", "kafka", "airflow", "big data",
    ],
    "methodologies": [
        "agile", "scrum", "kanban", "devops", "tdd", "bdd", "rest", "restful", "graphql",
        "microservices", "oop", "design patterns", "mvc", "unit testing",
        "integration testing", "api", "sdlc", "waterfall",
    ],
    "soft": [
        "communication", "teamwork", "leadership", "problem solving", "critical thinking",
        "time management", "collaboration", "adaptability", "creativity", "analytical",
        "attention to detail", "mentoring", "stakeholder management", "presentation",
    ],
}

# Flat skill set + the category each skill belongs to.
_SKILL_CATEGORY: dict[str, str] = {}
for _cat, _items in _SKILLS.items():
    for _s in _items:
        _SKILL_CATEGORY[_s] = _cat
_ALL_SKILLS = set(_SKILL_CATEGORY)

# Aliases → canonical skill, so "JS" and "react.js" don't read as misses.
_ALIASES: dict[str, str] = {
    "js": "javascript", "ts": "typescript", "py": "python", "golang": "go",
    "node.js": "node", "nodejs": "node", "reactjs": "react", "react.js": "react",
    "vue.js": "vue", "vuejs": "vue", "nextjs": "next.js", "postgres": "postgresql",
    "k8s": "kubernetes", "ml": "machine learning", "dl": "deep learning",
    "ai": "machine learning", "gcp": "google cloud", "ms sql": "sql server",
    "mssql": "sql server", "tf": "tensorflow", "cicd": "ci/cd", "c sharp": "c#",
    "dotnet": ".net", "rest api": "rest", "restful api": "restful",
    "unit tests": "unit testing", "oops": "oop",
}

# Standard résumé sections (canonical → recognised heading synonyms).
_SECTIONS: dict[str, list[str]] = {
    "summary": ["summary", "objective", "profile", "about", "professional summary", "career objective"],
    "experience": ["experience", "work experience", "employment", "professional experience", "work history"],
    "education": ["education", "academic", "qualifications", "academic background"],
    "skills": ["skills", "technical skills", "core competencies", "competencies", "technologies"],
    "projects": ["projects", "personal projects", "key projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses", "courses"],
}

_STOP = set(
    ("a an the and or but if then of to in on at by for with as is are was were be been "
     "being have has had do does did will would shall should can could may might must this "
     "that these those i you he she it we they your our their from not no into over under "
     "out up down off about which who whom whose when where why how all any both each more "
     "most other some such only own same so than too very just also able etc via per within").split()
)

# Generic JD boilerplate — phrases made only of these add no signal as "keywords".
_BOILERPLATE = set(
    ("requirements responsibilities qualifications hiring looking candidate role position "
     "job experience experienced strong excellent good knowledge familiarity proficiency "
     "ability skills skill work working years plus must required preferred desirable team "
     "company we are you will join including etc related field degree").split()
)

_WORD = re.compile(r"[a-z0-9][a-z0-9+#.\-]*")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
_URL = re.compile(r"https?://\S+|linkedin\.com/\S+|github\.com/\S+", re.I)


# ──────────────────────────────────────────────────────────────────────────────
# Text + structure extraction
# ──────────────────────────────────────────────────────────────────────────────

@dataclass
class CvDoc:
    text: str
    lines: list[str]
    is_pdf: bool
    page_count: int = 1
    # Parse-ability signals (PDF only; benign defaults for DOCX).
    multi_column: bool = False
    has_tables: bool = False
    image_only_ratio: float = 0.0      # share of pages that look image-based
    header_footer_contact: bool = False


def _extract_pdf(data: bytes) -> CvDoc:
    """CV text + ATS-relevant layout signals from a PDF."""
    import fitz

    from app.services.pdf_structure import extract_tables, harvest_spans

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        all_text: list[str] = []
        multi_column = False
        has_tables = False
        image_pages = 0
        header_footer_contact = False

        for page in doc:
            txt = page.get_text("text")
            all_text.append(txt)

            spans = harvest_spans(page)
            # Image-only page: almost no extractable text but the page has images.
            if len(txt.strip()) < 40 and page.get_images():
                image_pages += 1

            if extract_tables(page):
                has_tables = True

            # Multi-column: a wide vertical gutter through the page with substantial
            # text on both sides. Detected directly (not via the structure engine's
            # XY-cut, which deliberately suppresses column splits it reads as tables —
            # but for ATS both a 2-column layout AND a table are parse hazards).
            if _is_multi_column(spans, page.rect.width, page.rect.height):
                multi_column = True

            # Contact info living in the top/bottom margin (some ATS drop headers).
            h = page.rect.height
            for s in spans:
                if (s.y1 < h * 0.08 or s.y0 > h * 0.92) and (_EMAIL.search(s.text) or _PHONE.search(s.text)):
                    header_footer_contact = True

        text = "\n".join(all_text)
        return CvDoc(
            text=text,
            lines=[ln.strip() for ln in text.splitlines() if ln.strip()],
            is_pdf=True,
            page_count=doc.page_count,
            multi_column=multi_column,
            has_tables=has_tables,
            image_only_ratio=(image_pages / max(1, doc.page_count)),
            header_footer_contact=header_footer_contact,
        )
    finally:
        doc.close()


def _is_multi_column(spans: list, page_width: float, page_height: float) -> bool:
    """True if the page has a genuine two-column layout (an ATS parse hazard).

    Looks for a wide empty vertical band (gutter) through the central region of the
    page with substantial text on BOTH sides, each side spanning a real vertical
    extent and the two sides overlapping vertically. This catches the sidebar/CV
    template layouts ATS parsers read across and scramble — independent of the
    table/column disambiguation the structure engine does for reading order.
    """
    if len(spans) < 8:
        return False
    # Merge occupied x-intervals across all spans.
    merged: list[list[float]] = []
    for a, b in sorted((s.x0, s.x1) for s in spans):
        if merged and a <= merged[-1][1] + 2:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])

    min_gap = max(36.0, page_width * 0.06)
    for (a1, b1), (a2, b2) in zip(merged, merged[1:]):
        gap = a2 - b1
        mid = (b1 + a2) / 2
        if gap < min_gap or not (page_width * 0.2 < mid < page_width * 0.8):
            continue
        left = [s for s in spans if s.x1 <= b1]
        right = [s for s in spans if s.x0 >= a2]
        # A real two-column body has several lines down BOTH sides (≥4 each) — this
        # rejects a one-line header that merely has the name left and contact right.
        if len(left) < 4 or len(right) < 4:
            continue
        l_y0, l_y1 = min(s.y0 for s in left), max(s.y1 for s in left)
        r_y0, r_y1 = min(s.y0 for s in right), max(s.y1 for s in right)
        v_overlap = min(l_y1, r_y1) - max(l_y0, r_y0)
        if v_overlap > 30:   # the two columns genuinely run alongside each other
            return True
    return False


def _extract_docx(data: bytes) -> CvDoc:
    """CV text from a DOCX (DOCX is generally ATS-friendly — light checks only)."""
    import docx

    d = docx.Document(io.BytesIO(data))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    has_tables = len(d.tables) > 0
    # Pull table cell text too (so skills in a table still count).
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paras.append(cell.text.strip())
    text = "\n".join(paras)
    return CvDoc(text=text, lines=paras, is_pdf=False, has_tables=has_tables)


def extract_cv(data: bytes, filename: str) -> CvDoc:
    """Extract CV text + layout signals, dispatching on file type."""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        return _extract_docx(data)
    return _extract_pdf(data)


# ──────────────────────────────────────────────────────────────────────────────
# Tokenisation, normalisation, skill + keyphrase matching
# ──────────────────────────────────────────────────────────────────────────────

def _normalize(text: str) -> str:
    """Lowercase, apply aliases, collapse whitespace — the comparison surface."""
    t = " " + text.lower() + " "
    for alias, canon in _ALIASES.items():
        t = re.sub(rf"(?<![\w]){re.escape(alias)}(?![\w])", canon, t)
    return re.sub(r"\s+", " ", t)


def _find_skills(text: str) -> set[str]:
    """Skills from the taxonomy present in the (normalised) text.

    Multi-word skills are matched as phrases; single tokens on word boundaries. This
    avoids false hits like 'r' inside 'react' or 'go' inside 'good'.
    """
    norm = _normalize(text)
    found: set[str] = set()
    for skill in _ALL_SKILLS:
        # Word-boundary match; skills may contain ., +, #, - which \b mishandles,
        # so anchor on non-alphanumerics instead.
        pat = rf"(?<![a-z0-9]){re.escape(skill)}(?![a-z0-9])"
        if re.search(pat, norm):
            found.add(skill)
    return found


def _tokens(text: str, keep_stop: bool = False) -> list[str]:
    out = []
    for w in _WORD.findall(text.lower()):
        if len(w) < 2 and not w.isdigit():
            continue
        if not keep_stop and w in _STOP:
            continue
        out.append(w)
    return out


def rake_keyphrases(text: str, top_k: int = 25) -> list[tuple[str, float]]:
    """RAKE keyphrase extraction (Rose et al.).

    Split text into candidate phrases at stop words / punctuation, score each word by
    degree(word)/frequency(word) over the phrase co-occurrence graph, and a phrase's
    score is the sum of its word scores. Self-contained — needs no background corpus —
    which is why it suits a single JD. Returns the top phrases with scores.
    """
    # Candidate phrases: runs of content words between stop words / punctuation.
    phrases: list[list[str]] = []
    cur: list[str] = []
    for raw in re.split(r"[^\w+#.\-]+", text.lower()):
        w = raw.strip(".-")
        if not w:
            if cur:
                phrases.append(cur); cur = []
            continue
        if w in _STOP or (len(w) < 2 and not w.isdigit()):
            if cur:
                phrases.append(cur); cur = []
            continue
        cur.append(w)
    if cur:
        phrases.append(cur)

    freq: dict[str, int] = {}
    degree: dict[str, int] = {}
    for ph in phrases:
        deg = len(ph) - 1
        for w in ph:
            freq[w] = freq.get(w, 0) + 1
            degree[w] = degree.get(w, 0) + deg + 1
    word_score = {w: degree[w] / freq[w] for w in freq}

    scored: dict[str, float] = {}
    for ph in phrases:
        if len(ph) > 4:
            continue
        text_ph = " ".join(ph)
        scored[text_ph] = max(scored.get(text_ph, 0), sum(word_score[w] for w in ph))

    return sorted(scored.items(), key=lambda kv: kv[1], reverse=True)[:top_k]


def _tfidf_cosine(a: str, b: str) -> float:
    """TF-IDF cosine similarity between two documents (2-doc corpus)."""
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    from collections import Counter
    import math

    ca, cb = Counter(ta), Counter(tb)
    vocab = set(ca) | set(cb)
    # df over the 2 docs.
    df = {w: (1 if w in ca else 0) + (1 if w in cb else 0) for w in vocab}
    idf = {w: math.log(1 + 2 / df[w]) for w in vocab}
    va = {w: ca.get(w, 0) * idf[w] for w in vocab}
    vb = {w: cb.get(w, 0) * idf[w] for w in vocab}
    dot = sum(va[w] * vb[w] for w in vocab)
    na = math.sqrt(sum(v * v for v in va.values()))
    nb = math.sqrt(sum(v * v for v in vb.values()))
    return dot / (na * nb) if na and nb else 0.0


def _detect_sections(lines: list[str]) -> set[str]:
    """Which standard résumé sections the CV appears to contain."""
    found: set[str] = set()
    for ln in lines:
        low = ln.lower().strip(" :•-")
        if len(low) > 40:
            continue
        for canon, names in _SECTIONS.items():
            if any(low == nm or low.startswith(nm) for nm in names):
                found.add(canon)
    return found


# ──────────────────────────────────────────────────────────────────────────────
# Scoring + gap report
# ──────────────────────────────────────────────────────────────────────────────

@dataclass
class AtsReport:
    overall: int
    sub_scores: dict[str, int]
    matched_skills: list[str]
    missing_skills: list[str]
    matched_keywords: list[str]
    missing_keywords: list[str]
    format_issues: list[dict]
    sections_present: list[str]
    sections_missing: list[str]
    contact: dict
    stats: dict = field(default_factory=dict)


def _format_checks(cv: CvDoc) -> tuple[int, list[dict]]:
    """ATS parse-ability score (0-100) + the issues found, each with a fix."""
    issues: list[dict] = []
    score = 100

    if cv.is_pdf and cv.multi_column:
        score -= 25
        issues.append({
            "severity": "high",
            "title": "Multi-column layout",
            "detail": "ATS parsers read left-to-right across the whole page, so columns get interleaved and garbled.",
            "fix": "Use a single-column layout for the main content.",
        })
    if cv.has_tables:
        score -= 18
        issues.append({
            "severity": "high",
            "title": "Tables detected",
            "detail": "Many ATS flatten or drop tables, scrambling the cell order.",
            "fix": "Replace tables with plain text lines or simple bullet points.",
        })
    if cv.image_only_ratio > 0.3:
        score -= 30
        issues.append({
            "severity": "high",
            "title": "Image-based / scanned content",
            "detail": "Text inside images is invisible to an ATS — it sees an empty page.",
            "fix": "Export the CV as real selectable text, not a scan or screenshot.",
        })
    if cv.header_footer_contact:
        score -= 10
        issues.append({
            "severity": "medium",
            "title": "Contact info in header/footer",
            "detail": "Some ATS ignore page headers and footers, so your email/phone may be missed.",
            "fix": "Put your name and contact details in the body of the first page.",
        })
    if cv.page_count > 2:
        score -= 6
        issues.append({
            "severity": "low",
            "title": f"{cv.page_count} pages",
            "detail": "Long CVs dilute keyword density and may not be fully parsed.",
            "fix": "Aim for 1–2 pages unless you have extensive senior experience.",
        })
    return max(0, score), issues


def analyze_resume(cv_bytes: bytes, filename: str, jd_text: str) -> AtsReport:
    """Full deterministic ATS analysis of a CV against a job description."""
    cv = extract_cv(cv_bytes, filename)
    cv_text = cv.text or ""
    jd_text = jd_text or ""

    # 1) Skills present in JD vs CV.
    jd_skills = _find_skills(jd_text)
    cv_skills = _find_skills(cv_text)
    matched_skills = sorted(jd_skills & cv_skills, key=lambda s: (_SKILL_CATEGORY.get(s, ""), s))
    missing_skills = sorted(jd_skills - cv_skills, key=lambda s: (_SKILL_CATEGORY.get(s, ""), s))

    # 2) JD keyphrases (RAKE) vs CV — keyword coverage beyond the skills taxonomy.
    norm_cv = _normalize(cv_text)
    jd_phrases = rake_keyphrases(jd_text, top_k=30)
    matched_kw, missing_kw = [], []
    for phrase, _score in jd_phrases:
        words = phrase.split()
        # Skip phrases that are entirely generic JD boilerplate (no real signal).
        if all(w in _BOILERPLATE for w in words):
            continue
        canon = _normalize(phrase).strip()
        # A phrase counts as present if its head words appear in the CV.
        head = canon.split()
        present = canon in norm_cv or (len(head) >= 2 and all(
            re.search(rf"(?<![a-z0-9]){re.escape(w)}(?![a-z0-9])", norm_cv) for w in head
        ))
        (matched_kw if present else missing_kw).append(phrase)

    # 3) Sub-scores.
    skills_score = round(100 * len(matched_skills) / len(jd_skills)) if jd_skills else 100
    kw_total = len(matched_kw) + len(missing_kw)
    keyword_score = round(100 * len(matched_kw) / kw_total) if kw_total else 100
    similarity_score = round(100 * _tfidf_cosine(jd_text, cv_text))
    format_score, format_issues = _format_checks(cv)

    overall = round(
        0.38 * keyword_score + 0.27 * skills_score + 0.15 * similarity_score + 0.20 * format_score
    )

    # 4) Sections + contact completeness.
    sections = _detect_sections(cv.lines)
    sections_missing = [s for s in ("summary", "experience", "education", "skills") if s not in sections]
    contact = {
        "email": bool(_EMAIL.search(cv_text)),
        "phone": bool(_PHONE.search(cv_text)),
        "links": bool(_URL.search(cv_text)),
    }

    return AtsReport(
        overall=overall,
        sub_scores={
            "keywords": keyword_score,
            "skills": skills_score,
            "similarity": similarity_score,
            "format": format_score,
        },
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        matched_keywords=matched_kw[:15],
        missing_keywords=missing_kw[:15],
        format_issues=format_issues,
        sections_present=sorted(sections),
        sections_missing=sections_missing,
        contact=contact,
        stats={
            "cvWords": len(_tokens(cv_text, keep_stop=True)),
            "jdSkills": len(jd_skills),
            "pageCount": cv.page_count,
            "fileType": "pdf" if cv.is_pdf else "docx",
        },
    )


def report_to_dict(r: AtsReport) -> dict:
    """JSON-serialisable view of the report for the API."""
    return {
        "overall": r.overall,
        "subScores": r.sub_scores,
        "matchedSkills": r.matched_skills,
        "missingSkills": r.missing_skills,
        "matchedKeywords": r.matched_keywords,
        "missingKeywords": r.missing_keywords,
        "formatIssues": r.format_issues,
        "sectionsPresent": r.sections_present,
        "sectionsMissing": r.sections_missing,
        "contact": r.contact,
        "stats": r.stats,
    }
